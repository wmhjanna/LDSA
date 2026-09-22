from __future__ import annotations

import csv
import json
import math
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import numpy as np
import openpyxl
import pandas as pd
import xlsxwriter


ROOT = Path("@INPUT_ROOT@/data_sources")
RAW_ROOT = ROOT / "source_data"
ENRICHED_ROOT = ROOT / "predictive_models" / "pollutant_data" / "all_sites"
LAN_FILE = ROOT / "LAN" / "PNSD_data.xlsx"
OUT_ROOT = Path("outputs/site_data_pnsd_cm3_with_bins")
STAGE_ROOT = Path("work/site_data_with_bins_stage")

SITES = [
    "ANB_UB", "ATH_SUB", "ATH_UB", "BCN_UB", "BOS_UB", "BUD_UB",
    "DEU_UB", "DRE_TR", "DRE_UB", "HEL_TR", "HEL_UB", "HOH_RB",
    "IPR_RB", "LAN_UB", "LEC_UB", "LEI2_TR", "LEI2_UB", "LEI_TR",
    "LEI_UB", "LND_TR", "LND_UB", "MAD_UB", "MEL_RB", "MUL_SUB",
    "MUL_UB", "NEU_SUB", "PRA_SUB", "ROC_UB", "SCH_SUB", "STO_TR",
    "WAL_SUB", "ZUG_SUB", "ZUR_UB",
]

BASE_HEADERS = [
    "Site", "Type", "UTC start", "UTC end",
    "BC (ng/m³)", "PM10 (µg/m³)", "PM2.5 (µg/m³)", "PM1 (µg/m³)",
    "SO2 (µg/m³)", "NO (µg/m³)", "NO2 (µg/m³)", "O3 (µg/m³)",
    "CO (mg/m³)", "TEMP (°C)", "RH (%)", "PRESS (hPa)",
    "RAD (W/m²)", "WS (m/s)", "WD (°)", "RAIN (mm)",
]

MODE_HEADERS = [
    "Nucleation (#/cm³)", "Aitken (#/cm³)", "Accumulation (#/cm³)",
    "N25-800 (#/cm³)", "Ntotal (#/cm³)",
]

TYPE_MAP = {
    "UB": "Urban background",
    "TR": "Traffic",
    "SUB": "Suburban background",
    "RB": "Rural background",
}


def norm_header(value) -> str:
    text = "" if value is None else str(value)
    text = text.replace("₂", "2").replace("₁", "1").replace("₀", "0")
    text = text.replace("³", "3").replace("μ", "u").replace("µ", "u")
    return re.sub(r"[^A-Z0-9]+", "", text.upper())


def output_header_key(header) -> str | None:
    h = norm_header(header)
    if h in {"SITE", "SITENAME"}: return "Site"
    if h in {"TYPE", "STATIONTYPE"}: return "Type"
    if ("UTC" in h or "LOCALTIME" in h) and "START" in h: return "UTC start"
    if ("UTC" in h or "LOCALTIME" in h) and "END" in h: return "UTC end"
    if h.startswith("BC"): return "BC (ng/m³)"
    if h.startswith("PM10"): return "PM10 (µg/m³)"
    if h.startswith("PM25"): return "PM2.5 (µg/m³)"
    if h.startswith("PM1") and not h.startswith("PM10"): return "PM1 (µg/m³)"
    if h.startswith("SO2"): return "SO2 (µg/m³)"
    if h.startswith("NO2"): return "NO2 (µg/m³)"
    if h == "NO" or h.startswith("NOUG") or h.startswith("NOMG"): return "NO (µg/m³)"
    if h.startswith("O3"): return "O3 (µg/m³)"
    if h.startswith("CO"): return "CO (mg/m³)"
    if h.startswith("TEMP"): return "TEMP (°C)"
    if h == "RH" or h.startswith("RH"): return "RH (%)"
    if h.startswith("PRESS"): return "PRESS (hPa)"
    if h.startswith("RAD"): return "RAD (W/m²)"
    if h == "WS" or h.startswith("WSMS"): return "WS (m/s)"
    if h == "WD" or h.startswith("WDDEG"): return "WD (°)"
    if h.startswith("RAIN"): return "RAIN (mm)"
    if h == "NUCLEATION": return "Nucleation (#/cm³)"
    if h == "AITKEN": return "Aitken (#/cm³)"
    if h == "ACCUMULATION": return "Accumulation (#/cm³)"
    if h.startswith("N25800"): return "N25-800 (#/cm³)"
    if h == "NTOTAL": return "Ntotal (#/cm³)"
    return None


def clean_value(value):
    if value is None:
        return None
    if isinstance(value, str) and value.strip().upper() in {
        "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!",
    }:
        return None
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    return value


def write_row(ws, rownum: int, values):
    cleaned = []
    for value in values:
        value = clean_value(value)
        if isinstance(value, pd.Timestamp):
            value = value.to_pydatetime()
        elif isinstance(value, np.integer):
            value = int(value)
        elif isinstance(value, np.floating):
            value = float(value)
        cleaned.append(value)
    ws.write_row(rownum, 0, cleaned)


def create_workbook(path: Path, headers):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = xlsxwriter.Workbook(path, {"constant_memory": True, "nan_inf_to_errors": True})
    ws = wb.add_worksheet("Data")
    header_fmt = wb.add_format({
        "bold": True, "font_color": "#FFFFFF", "bg_color": "#1F4E78",
        "border": 1, "border_color": "#D9E2F3", "align": "center", "valign": "vcenter",
    })
    date_fmt = wb.add_format({"num_format": "yyyy-mm-dd hh:mm:ss"})
    num_fmt = wb.add_format({"num_format": "0.000"})
    ws.write_row(0, 0, headers, header_fmt)
    ws.freeze_panes(1, 4)
    ws.autofilter(0, 0, 0, len(headers) - 1)
    ws.set_row(0, 26)
    ws.set_column(0, 0, 18)
    ws.set_column(1, 1, 22)
    ws.set_column(2, 3, 20, date_fmt)
    ws.set_column(4, len(headers) - 1, 15, num_fmt)
    first_mode_col = len(headers) - len(MODE_HEADERS)
    ws.set_column(first_mode_col, len(headers) - 1, 20, num_fmt)
    return wb, ws, date_fmt, num_fmt


def format_dp(dp: float) -> str:
    text = f"{dp:.3f}".rstrip("0").rstrip(".")
    return f"D_{text} nm (#/cm³)"


def diameter_columns(headers):
    found = []
    for idx, header in enumerate(headers):
        if isinstance(header, (int, float)) and 1 <= float(header) <= 2000:
            found.append((idx, float(header)))
    return sorted(found, key=lambda item: item[1])


def delta_log10(dps):
    dps = np.asarray(dps, dtype=float)
    delta = np.diff(np.log10(dps))
    return np.append(delta, delta[-1])


def pick_enriched_sheet(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    best = None
    best_score = -1
    for ws in wb.worksheets:
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        score = sum(output_header_key(h) is not None for h in headers)
        if score > best_score:
            best, best_score = ws, score
    return wb, best


def build_from_enriched(site: str, source: Path, stage_path: Path):
    in_wb, in_ws = pick_enriched_sheet(source)
    header_cells = next(in_ws.iter_rows(min_row=1, max_row=1))
    index_by_key = {}
    for idx, cell in enumerate(header_cells):
        key = output_header_key(cell.value)
        if key and key not in index_by_key:
            index_by_key[key] = idx

    source_headers = [cell.value for cell in header_cells]
    dp_columns = diameter_columns(source_headers)
    dps = [dp for _, dp in dp_columns]
    weights = delta_log10(dps)
    headers = BASE_HEADERS + [format_dp(dp) for dp in dps] + MODE_HEADERS
    out_wb, out_ws, date_fmt, num_fmt = create_workbook(stage_path, headers)
    out_row = 1
    blank_run = 0
    site_idx = index_by_key.get("Site")
    type_idx = index_by_key.get("Type")
    start_idx = index_by_key.get("UTC start")

    for source_row in in_ws.iter_rows(min_row=2, values_only=True):
        start = source_row[start_idx] if start_idx is not None and start_idx < len(source_row) else None
        if start is None:
            blank_run += 1
            if blank_run >= 2000 and out_row > 1:
                break
            continue
        blank_run = 0
        base_values = []
        for header in BASE_HEADERS:
            idx = index_by_key.get(header)
            base_values.append(source_row[idx] if idx is not None and idx < len(source_row) else None)
        bin_values = []
        for (idx, _), weight in zip(dp_columns, weights):
            raw = source_row[idx] if idx < len(source_row) else None
            try:
                bin_values.append(float(raw) * float(weight) if raw is not None else None)
            except (TypeError, ValueError):
                bin_values.append(None)
        mode_values = []
        for header in MODE_HEADERS:
            idx = index_by_key.get(header)
            mode_values.append(source_row[idx] if idx is not None and idx < len(source_row) else None)
        values = base_values + bin_values + mode_values
        if site == "ROC_UB":
            # ROC source timestamps are explicitly local Rochester time; normalize them to UTC.
            eastern = ZoneInfo("America/New_York")
            for time_pos in (2, 3):
                if isinstance(values[time_pos], datetime):
                    values[time_pos] = values[time_pos].replace(tzinfo=eastern).astimezone(timezone.utc).replace(tzinfo=None)
        if not values[0]: values[0] = source_row[site_idx] if site_idx is not None else site
        if not values[1]: values[1] = source_row[type_idx] if type_idx is not None else TYPE_MAP[site.rsplit("_", 1)[-1]]
        # Use the final integrated values already present in the source workbook.
        mode_start = len(BASE_HEADERS) + len(bin_values)
        nuc, ait, acc, n25800, ntotal = values[mode_start:mode_start + 5]
        if n25800 is None and ait is not None and acc is not None:
            values[mode_start + 3] = clean_value(ait) + clean_value(acc)
        if ntotal is None and nuc is not None and values[mode_start + 3] is not None:
            values[mode_start + 4] = clean_value(nuc) + clean_value(values[mode_start + 3])
        if ait is None and ntotal is not None and nuc is not None and acc is not None:
            values[mode_start + 1] = clean_value(ntotal) - clean_value(nuc) - clean_value(acc)
            values[mode_start + 3] = clean_value(values[mode_start + 1]) + clean_value(acc)
        write_row(out_ws, out_row, values)
        out_row += 1
    out_ws.autofilter(0, 0, max(1, out_row - 1), len(headers) - 1)
    out_wb.close()
    in_wb.close()
    return out_row - 1, len(dps)


def extract_diameter(header) -> float | None:
    m = re.search(r"[-+]?\d+(?:\.\d+)?", str(header))
    return float(m.group()) if m else None


def calc_pnsd(frame: pd.DataFrame):
    diameters, columns = [], []
    for col in frame.columns:
        dp = extract_diameter(col)
        if dp is not None and not norm_header(col).startswith(("TEMP", "PRESS")):
            diameters.append(dp)
            columns.append(col)
    if len(columns) < 2:
        raise ValueError("No usable PNSD diameter columns")
    order = np.argsort(diameters)
    dps = np.asarray(diameters, dtype=float)[order]
    columns = [columns[i] for i in order]
    matrix = frame[columns].apply(pd.to_numeric, errors="coerce").to_numpy(dtype=float)
    delta = delta_log10(dps)
    per_bin = matrix * delta

    def integrate(low, high):
        mask = (dps >= low) & (dps < high)
        if not mask.any():
            return np.full(len(frame), np.nan)
        block = matrix[:, mask]
        weights = delta[mask]
        valid = np.isfinite(block)
        result = np.nansum(block * weights, axis=1)
        result[~valid.any(axis=1)] = np.nan
        return result

    nuc = integrate(10.0, 25.0)
    ait = integrate(25.0, 100.0)
    acc = integrate(100.0, 800.000001)
    return dps, per_bin, (nuc, ait, acc, ait + acc, nuc + ait + acc)


def read_raw_site(site: str):
    if site == "LAN_UB":
        book = pd.ExcelFile(LAN_FILE)
        frames = []
        for sheet in book.sheet_names:
            df = book.parse(sheet)
            if len(df.columns) >= 10:
                frames.append(df)
        return frames, [f"{LAN_FILE}#{s}" for s in book.sheet_names]
    if site == "PRA_SUB":
        paths = sorted(RAW_ROOT.glob("PRA_SUB_*/*.xlsx"))
    else:
        paths = [RAW_ROOT / site / f"{site}.xlsx"]
    frames = [pd.read_excel(p, sheet_name=0) for p in paths]
    return frames, [str(p) for p in paths]


def build_from_raw(site: str, stage_path: Path):
    frames, sources = read_raw_site(site)
    calculations = [calc_pnsd(frame) for frame in frames]
    all_dps = sorted({round(float(dp), 6) for dps, _, _ in calculations for dp in dps})
    headers = BASE_HEADERS + [format_dp(dp) for dp in all_dps] + MODE_HEADERS
    dp_position = {dp: idx for idx, dp in enumerate(all_dps)}
    out_wb, out_ws, date_fmt, num_fmt = create_workbook(stage_path, headers)
    out_row = 1
    last_interval = timedelta(hours=1)
    for frame, (dps, per_bin, modes) in zip(frames, calculations):
        time_col = next((c for c in frame.columns if norm_header(c) in {"TIME", "DATETIME", "DATE"}), frame.columns[0])
        times = pd.to_datetime(frame[time_col], errors="coerce")
        valid_times = times.dropna().sort_values()
        diffs = valid_times.diff().dropna()
        diffs = diffs[(diffs > pd.Timedelta(0)) & (diffs <= pd.Timedelta(days=1))]
        if not diffs.empty:
            last_interval = diffs.median().to_pytimedelta()
        frame_dp_positions = [dp_position[round(float(dp), 6)] for dp in dps]
        temp_col = next((c for c in frame.columns if norm_header(c).startswith("TEMP")), None)
        press_col = next((c for c in frame.columns if norm_header(c).startswith("PRESS")), None)
        temp_values = pd.to_numeric(frame[temp_col], errors="coerce") if temp_col else None
        press_values = pd.to_numeric(frame[press_col], errors="coerce") if press_col else None
        for i, start in enumerate(times):
            if pd.isna(start):
                continue
            end = start + last_interval - timedelta(seconds=1)
            base = [None] * len(BASE_HEADERS)
            bins = [None] * len(all_dps)
            mode_values = [None] * len(MODE_HEADERS)
            base[0] = site
            base[1] = TYPE_MAP[site.rsplit("_", 1)[-1]]
            base[2] = start
            base[3] = end
            if temp_values is not None:
                temp = temp_values.iloc[i]
                base[13] = (temp - 273.15) if pd.notna(temp) and temp > 150 else temp
            if press_values is not None:
                base[15] = press_values.iloc[i]
            for source_pos, output_pos in enumerate(frame_dp_positions):
                bins[output_pos] = per_bin[i, source_pos]
            for j, arr in enumerate(modes):
                mode_values[j] = arr[i]
            write_row(out_ws, out_row, base + bins + mode_values)
            out_row += 1
    out_ws.autofilter(0, 0, max(1, out_row - 1), len(headers) - 1)
    out_wb.close()
    return out_row - 1, sources, len(all_dps)


def main():
    STAGE_ROOT.mkdir(parents=True, exist_ok=True)
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    manifest = []
    for number, site in enumerate(SITES, 1):
        stage_path = STAGE_ROOT / site / f"{site}_standardized.xlsx"
        enriched = ENRICHED_ROOT / f"{site}.xlsx"
        if enriched.exists():
            rows, size_bins = build_from_enriched(site, enriched, stage_path)
            method = "Existing final #/cm³ columns retained"
            sources = [str(enriched)]
        else:
            rows, sources, size_bins = build_from_raw(site, stage_path)
            method = "Integrated from PNSD: dN/dlogDp × Δlog10Dp"
        manifest.append({
            "site": site, "rows": rows, "method": method,
            "size_bins": size_bins, "source": " | ".join(sources),
            "stage_file": str(stage_path),
        })
        print(f"[{number:02d}/{len(SITES)}] {site}: {rows:,} rows — {method}", flush=True)
    with (STAGE_ROOT / "manifest.json").open("w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)
    with (STAGE_ROOT / "manifest.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=manifest[0].keys())
        writer.writeheader(); writer.writerows(manifest)


if __name__ == "__main__":
    main()
