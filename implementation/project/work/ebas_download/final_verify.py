from pathlib import Path
import json,collections,zipfile
import openpyxl
W=Path('work/ebas_download');O=Path('@DELIVERY_ROOT@/EBAS_observations')
plan=json.loads((W/'selection_plan.json').read_text());checks=json.loads((W/'metadata_checks.json').read_text())
if (W/'additional_metadata_checks.json').exists():checks+=json.loads((W/'additional_metadata_checks.json').read_text())
if (W/'replacement_metadata_checks.json').exists():checks+=json.loads((W/'replacement_metadata_checks.json').read_text())
expected={r['name'] for k in ['PNSD','direct_pollutants'] for r in plan[k]}|{r['file'] for r in checks if r.get('variables')}
d={r['file']:r for r in map(json.loads,(O/'download_log.jsonl').read_text().splitlines())}
replacements=json.loads((O/'source_replacements.json').read_text()) if (O/'source_replacements.json').exists() else []
for r in replacements:expected.discard(r['old_file']);expected.add(r['new_file']);d.pop(r['old_file'],None)
a={r['file']:r for r in map(json.loads,(O/'verified_metadata.jsonl').read_text().splitlines())}
assert set(d)==expected,(len(d),len(expected),expected-set(d))
assert all(v['status']=='downloaded' for v in d.values()),[x for x in d.values() if x['status']!='downloaded']
assert set(a)==expected,(expected-set(a),set(a)-expected)
assert all(v['status']=='verified' for v in a.values()),[x for x in a.values() if x['status']!='verified']
for r in d.values():assert (O/r['relative_path']).stat().st_size==r['bytes']
(O/'download_manifest.json').write_text(json.dumps(list(d.values()),ensure_ascii=False,indent=2))
(O/'verified_metadata.json').write_text(json.dumps(list(a.values()),ensure_ascii=False,indent=2))
summary={'metadata_candidates_inspected':len(checks),'metadata_status':dict(collections.Counter(r['status'] for r in checks)),'extra_matching_files':sum(bool(r.get('variables')) for r in checks),'expected_files':len(expected),'downloaded_files':len(d),'verified_files':len(a),'total_bytes':sum(r['bytes'] for r in d.values()),'category_counts':dict(collections.Counter(r['category'] for r in d.values())),'available_station_counts':{m:len({r['station_code'] for r in a.values() if m in r.get('available_metrics',[])}) for m in ['PNSD','PM2.5','PM10','PM1','NO2','NO','SO2','O3','CO','BC']}}
(O/'data_validation.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
(O/'SHA256SUMS.txt').write_text(''.join(f"{r['sha256']}  {r['relative_path']}\n" for r in sorted(d.values(),key=lambda r:r['relative_path'])))
wb=openpyxl.load_workbook(O/'EBAS_station_inventory.xlsx',read_only=True,data_only=True)
assert sum(1 for r in wb['Stations'].iter_rows(values_only=True) if any(v is not None for v in r))==66
assert sum(1 for r in wb['Configurations'].iter_rows(values_only=True) if any(v is not None for v in r))==651
assert sum(1 for r in wb['Source_files'].iter_rows(values_only=True) if any(v is not None for v in r))==len(expected)+1
assert sum(1 for r in wb['Analysis_inputs'].iter_rows(values_only=True) if any(v is not None for v in r))==35
assert len({r[0] for r in wb['Stations'].iter_rows(min_row=2,values_only=True)})==65
assert set(r[2] for r in wb['Source_files'].iter_rows(min_row=2,values_only=True))==expected
print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)
