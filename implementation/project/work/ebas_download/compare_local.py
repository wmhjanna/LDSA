from pathlib import Path
import json,re,collections,datetime
import openpyxl,pandas as pd
W=Path('work/ebas_download');stations=json.loads((W/'stations.json').read_text())
paths=(W/'local_paths.txt').read_text().splitlines()
location=Path('@DELIVERY_ROOT@/station_locations.xlsx')
loc=pd.read_excel(location)
alias={'IPR_RB':'IT0004R','ATH_SUB':'GR0100B','PAR_SUB':'FR0020R','PRA_SUB':'CZ0004B','LEI_TR':'DE0067K','BCN_UB':'ES0019U','DRE_UB':'DE0064B','HEL_UB':'FI0038U','LEC_UB':'IT0015U','LEI_UB':'DE0055B','LND_UB':'GB0021U','MAD_UB':'ES0021U','MUL_UB':'DE0069B'}
mapping=collections.defaultdict(list)
for r in loc.to_dict('records'):
 code=r['ID'] if re.fullmatch(r'[A-Z]{2}\d{4}[A-Z]',str(r['ID'])) else alias.get(r['Acronym'])
 if code:mapping[code].append(r['Acronym'])
codehits={};result=[]
for s in stations:
 code=s['code'];hits=[p for p in paths if re.search('/'+code+r'\.',p) and Path(p).suffix.lower() in ['.xlsx','.xls','.nc','.nas'] and any('.'+inst+'.' in p for inst in ['smps','dmps','tdmps','tsmps','aps','optical_particle_size_spectrometer','optical_particle_counter'])]
 mobility=[p for p in hits if any('.'+i+'.' in p for i in ['smps','dmps','tdmps','tsmps'])]
 evidence=sorted(mobility or hits,key=lambda p:(0 if '/LDSA/data/' in p else 1,len(p)))
 if not evidence and mapping.get(code):
  evidence=[f'@INPUT_ROOT@/data_sources/source_data/{mapping[code][0]}/{mapping[code][0]}.xlsx']
 codehits[code]=hits
 info={**s,'current_analysis_sites':mapping.get(code,[]),'local_station_code_files':len(hits),'local_mobility_files':len(mobility),'local_evidence':evidence[0] if evidence else '', 'local_status':'included_in_study' if mapping.get(code) else 'local_MPSS_available' if mobility else 'other_local_instrument_available' if hits else 'PNSD_not_available_locally'}
 if evidence:
  p=Path(evidence[0]);assert p.exists(),p
  if p.suffix=='.xlsx':
   wb=openpyxl.load_workbook(p,read_only=True,data_only=True);sh=wb.worksheets[0];head=list(next(sh.iter_rows(max_row=1,values_only=True)));sample=list(sh.iter_rows(min_row=2,max_row=3,values_only=True));wb.close()
   info['sample_sheet']=sh.title;info['sample_rows']=sh.max_row;info['sample_columns']=sh.max_column;info['sample_header']=[str(v) for v in head[:10]]
   info['sample_first_time']=str(sample[0][0]) if sample else ''
   dia=[]
   for h in head:
    m=re.fullmatch(r'(?:D_|particle_number_size_distribution_)?([0-9]+(?:\.[0-9]+)?)',str(h))
    if m:dia.append(float(m[1]))
   info['sample_particle_channels']=len(dia);info['sample_min_nm']=min(dia) if dia else None;info['sample_max_nm']=max(dia) if dia else None
 notes=[]
 if code=='DE0069B':notes.append('Provider station: Mülheim-Styrum (DE0069B).')
 if code in ['FI0038U','IT0015U']:notes.append('Match observations using the provider station code and native instrument metadata.')
 if code=='IT0004R':notes.append('Provider and study station classifications differ; use the study classification RB for IPR_RB.')
 info['matching_notes']=' '.join(notes);result.append(info)
(W/'local_comparison.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
(W/'local_code_files.json').write_text(json.dumps(codehits,ensure_ascii=False,indent=2))
print(collections.Counter(x['local_status'] for x in result));print('CURRENT_MAPPED_CODES',sum(bool(x['current_analysis_sites']) for x in result));print('NEW_LOCAL',[x['name'] for x in result if x['local_status']=='PNSD_not_available_locally'])
