from __future__ import annotations

import re
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from matplotlib.ticker import MaxNLocator
from openpyxl import load_workbook


SOURCE = Path(
    "@ATTACHMENT_ROOT@/"
    "xwechat_files/wxid_u2wh4eh9sunh22_ea96/temp/drag/"
    "Fig3_Origin_ready_source_data(1).xlsx"
)
OUTPUT_DIR = Path(
    "@PROJECT_ROOT@/outputs/fig3_python_redraw"
)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


mpl.rcParams.update(
    {
        "font.family": "Times New Roman",
        "font.serif": ["Times New Roman"],
        "font.weight": "bold",
        "mathtext.fontset": "custom",
        "mathtext.rm": "Times New Roman",
        "mathtext.it": "Times New Roman:italic",
        "mathtext.bf": "Times New Roman:bold",
        "mathtext.default": "bf",
        "axes.labelweight": "bold",
        "axes.titleweight": "bold",
        "axes.linewidth": 1.25,
        "xtick.major.width": 1.15,
        "ytick.major.width": 1.15,
        "xtick.major.size": 4.5,
        "ytick.major.size": 4.5,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
        "svg.fonttype": "none",
    }
)


TYPE_STYLE = {
    "UB": {"color": "#0072B2", "marker": "o", "label": "UB"},
    "SUB": {"color": "#7B3294", "marker": "^", "label": "SUB"},
    "TR": {"color": "#D55E00", "marker": "s", "label": "TR"},
    "RB": {"color": "#009E73", "marker": "D", "label": "RB"},
}

METRIC_DISPLAY = {
    "PM2.5": r"$\mathbf{PM}_{\mathbf{2.5}}$",
    "PM10": r"$\mathbf{PM}_{\mathbf{10}}$",
    "BC": "BC",
    "NO2": r"$\mathbf{NO}_{\mathbf{2}}$",
    "SO2": r"$\mathbf{SO}_{\mathbf{2}}$",
}


def read_sheet(wb, sheet_name):
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    return [dict(zip(headers, row)) for row in rows[1:]]


def short_site(site):
    return re.sub(r"_(UB|SUB|TR|RB)$", "", str(site))


def style_axes(ax):
    ax.grid(True, color="#D9D9D9", linestyle=(0, (1.4, 2.2)), linewidth=0.8, zorder=0)
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_linewidth(1.25)
        spine.set_color("#222222")
    ax.tick_params(labelsize=10.5, pad=4)
    for tick in [*ax.get_xticklabels(), *ax.get_yticklabels()]:
        tick.set_fontweight("bold")


def panel_heading(ax, letter, title):
    ax.set_title(title, loc="left", fontsize=17, fontweight="bold", pad=8)
    ax.text(
        -0.075,
        1.025,
        letter,
        transform=ax.transAxes,
        ha="right",
        va="bottom",
        fontsize=18,
        fontweight="bold",
    )


def place_site_labels(fig, ax, point_records, stats_artist):
    """Place all site labels with short leader lines and basic collision avoidance."""
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    axes_box = ax.get_window_extent(renderer)
    occupied = []
    if stats_artist is not None:
        occupied.append(stats_artist.get_window_extent(renderer).expanded(1.08, 1.12))

    point_boxes = []
    for rec in point_records:
        px, py = ax.transData.transform((rec["x"], rec["y"]))
        point_boxes.append(mpl.transforms.Bbox.from_bounds(px - 6, py - 6, 12, 12))

    candidates = [
        (7, 6),
        (7, -7),
        (-7, 6),
        (-7, -7),
        (13, 0),
        (-13, 0),
        (13, 10),
        (-13, 10),
        (13, -11),
        (-13, -11),
        (20, 5),
        (-20, 5),
        (20, -7),
        (-20, -7),
    ]

    order = sorted(
        range(len(point_records)),
        key=lambda i: (-point_records[i]["y"], point_records[i]["x"]),
    )

    for idx in order:
        rec = point_records[idx]
        accepted = None
        for dx, dy in candidates:
            ann = ax.annotate(
                rec["label"],
                xy=(rec["x"], rec["y"]),
                xytext=(dx, dy),
                textcoords="offset points",
                ha="left" if dx >= 0 else "right",
                va="bottom" if dy >= 0 else "top",
                fontsize=7.2,
                fontweight="bold",
                color="#222222",
                bbox=dict(facecolor="white", edgecolor="none", alpha=0.82, pad=0.08),
                arrowprops=dict(
                    arrowstyle="-",
                    color="#777777",
                    linewidth=0.45,
                    shrinkA=1.5,
                    shrinkB=4.0,
                ),
                zorder=6,
                clip_on=True,
            )
            fig.canvas.draw()
            bbox = ann.get_window_extent(renderer)
            inside = (
                bbox.x0 >= axes_box.x0 + 2
                and bbox.x1 <= axes_box.x1 - 2
                and bbox.y0 >= axes_box.y0 + 2
                and bbox.y1 <= axes_box.y1 - 2
            )
            label_collision = any(bbox.overlaps(other) for other in occupied)
            point_collision = any(
                bbox.overlaps(pb) for j, pb in enumerate(point_boxes) if j != idx
            )
            if inside and not label_collision and not point_collision:
                accepted = ann
                occupied.append(bbox.expanded(1.04, 1.10))
                break
            ann.remove()

        if accepted is None:
            dx, dy = (8, 8)
            ann = ax.annotate(
                rec["label"],
                xy=(rec["x"], rec["y"]),
                xytext=(dx, dy),
                textcoords="offset points",
                ha="left",
                va="bottom",
                fontsize=6.8,
                fontweight="bold",
                color="#222222",
                bbox=dict(facecolor="white", edgecolor="none", alpha=0.84, pad=0.06),
                arrowprops=dict(
                    arrowstyle="-",
                    color="#777777",
                    linewidth=0.4,
                    shrinkA=1.5,
                    shrinkB=4.0,
                ),
                zorder=6,
                clip_on=True,
            )
            fig.canvas.draw()
            occupied.append(ann.get_window_extent(renderer).expanded(1.04, 1.08))


wb = load_workbook(SOURCE, data_only=True, read_only=True)
panel_sheets = {
    "a": ("Fig3a_PM2.5", "PM2.5"),
    "b": ("Fig3b_PM10", "PM10"),
    "c": ("Fig3c_BC", "BC"),
    "d": ("Fig3d_NO2", "NO2"),
}

fig = plt.figure(figsize=(15.8, 18.4))
grid = fig.add_gridspec(
    3,
    2,
    height_ratios=[1.0, 1.0, 1.0],
    left=0.085,
    right=0.985,
    bottom=0.085,
    top=0.985,
    wspace=0.16,
    hspace=0.23,
)

axes = {
    "a": fig.add_subplot(grid[0, 0]),
    "b": fig.add_subplot(grid[0, 1]),
    "c": fig.add_subplot(grid[1, 0]),
    "d": fig.add_subplot(grid[1, 1]),
    "e": fig.add_subplot(grid[2, 0]),
    "f": fig.add_subplot(grid[2, 1]),
}


# Panels a-d: exact-overlap slopes.
for letter, (sheet_name, metric) in panel_sheets.items():
    ax = axes[letter]
    rows = read_sheet(wb, sheet_name)
    xs = np.array([float(row["Pollutant slope (% yr^-1) [X]"]) for row in rows])
    ys = np.array([float(row["LDSA slope (% yr^-1) [Y]"]) for row in rows])
    lo = min(float(xs.min()), float(ys.min()))
    hi = max(float(xs.max()), float(ys.max()))
    span = max(hi - lo, 1.0)
    pad = span * 0.10
    lo -= pad
    hi += pad

    ax.plot([lo, hi], [lo, hi], linestyle="--", color="#777777", linewidth=1.35, zorder=1)
    ax.axhline(0, color="#AAAAAA", linewidth=0.9, zorder=1)
    ax.axvline(0, color="#AAAAAA", linewidth=0.9, zorder=1)

    point_records = []
    for row in rows:
        st = str(row["Station type"])
        style = TYPE_STYLE[st]
        x = float(row["Pollutant slope (% yr^-1) [X]"])
        y = float(row["LDSA slope (% yr^-1) [Y]"])
        ax.scatter(
            x,
            y,
            s=68,
            marker=style["marker"],
            color=style["color"],
            edgecolors="white",
            linewidths=0.8,
            zorder=4,
        )
        point_records.append({"x": x, "y": y, "label": short_site(row["Site"])})

    n = len(rows)
    r = float(np.corrcoef(xs, ys)[0, 1]) if n > 1 else float("nan")
    stats_artist = ax.text(
        0.035,
        0.965,
        f"n = {n}\nr = {r:.2f}",
        transform=ax.transAxes,
        ha="left",
        va="top",
        fontsize=11.5,
        fontweight="bold",
        linespacing=1.18,
        bbox=dict(facecolor="white", edgecolor="none", alpha=0.94, pad=1.5),
        zorder=8,
    )

    ax.set_xlim(lo, hi)
    ax.set_ylim(lo, hi)
    # Let the four scatter-panel frames fill the same GridSpec cells as the
    # panels below, so all left/right and top/bottom frame edges are aligned.
    ax.set_aspect("auto")
    ax.xaxis.set_major_locator(MaxNLocator(nbins=6))
    ax.yaxis.set_major_locator(MaxNLocator(nbins=6))
    ax.set_xlabel(
        f"{METRIC_DISPLAY[metric]} slope (% yr$^{{-1}}$)",
        fontsize=13.5,
        fontweight="bold",
        labelpad=7,
    )
    ax.set_ylabel(
        r"LDSA slope (% yr$^{-1}$)",
        fontsize=13.5,
        fontweight="bold",
        labelpad=7,
    )
    panel_heading(ax, letter, METRIC_DISPLAY[metric])
    style_axes(ax)
    place_site_labels(fig, ax, point_records, stats_artist)


# Panel e: trend-classification agreement.
ax = axes["e"]
rows_e = read_sheet(wb, "Fig3e_plot_data")
metric_order_e = ["PM2.5", "PM10", "BC", "NO2"]
rows_e = sorted(rows_e, key=lambda r: metric_order_e.index(r["Pollutant"]))
y = np.arange(len(rows_e))
segments = [
    ("Same classification (%)", "Same classification", "#BDBDBD"),
    ("LDSA decline only (%)", "LDSA decline only", "#56B4E9"),
    ("Pollutant decline only (%)", "Pollutant decline only", "#E69F00"),
    ("Other mismatch (%)", "Other mismatch", "#CC79A7"),
]
left = np.zeros(len(rows_e))
for col, label, color in segments:
    vals = np.array([float(r[col]) for r in rows_e])
    ax.barh(y, vals, left=left, color=color, edgecolor="white", linewidth=0.8, height=0.62)
    for yi, start, val in zip(y, left, vals):
        if val >= 10:
            ax.text(
                start + val / 2,
                yi,
                f"{val:.0f}%",
                ha="center",
                va="center",
                fontsize=10,
                fontweight="bold",
            )
    left += vals

ax.set_yticks(y)
ax.set_yticklabels([METRIC_DISPLAY[r["Pollutant"]] for r in rows_e], fontsize=12, fontweight="bold")
ax.invert_yaxis()
ax.set_xlim(0, 100)
ax.set_xlabel("Share of exact-overlap site pairs (%)", fontsize=13.5, fontweight="bold", labelpad=8)
ax.xaxis.set_major_locator(MaxNLocator(nbins=6))
panel_heading(ax, "e", "Trend classification agreement")
style_axes(ax)
ax.grid(axis="x", color="#D9D9D9", linestyle=(0, (1.4, 2.2)), linewidth=0.8)
ax.grid(axis="y", visible=False)
handles_e = [Patch(facecolor=color, edgecolor="none", label=label) for _, label, color in segments]
leg_e = ax.legend(
    handles=handles_e,
    loc="upper center",
    bbox_to_anchor=(0.5, -0.19),
    ncol=2,
    frameon=False,
    fontsize=10.5,
    columnspacing=1.6,
    handlelength=1.8,
)
for text in leg_e.get_texts():
    text.set_fontweight("bold")


# Panel f: metric-specific transition timing.
ax = axes["f"]
rows_f = read_sheet(wb, "Fig3f_offsets_raw")
summary_f = {r["Pollutant"]: r for r in read_sheet(wb, "Fig3f_summary")}
metric_order_f = ["PM2.5", "PM10", "BC", "NO2", "SO2"]
positions = np.arange(len(metric_order_f))
data_f = [
    [float(r["Offset Δt (months) [Y]"]) for r in rows_f if r["Pollutant"] == metric]
    for metric in metric_order_f
]

ax.axhspan(-12, 12, color="#E8E8E8", zorder=0)
box = ax.boxplot(
    data_f,
    positions=positions,
    widths=0.48,
    patch_artist=True,
    showfliers=False,
    medianprops=dict(color="#111111", linewidth=2.0),
    boxprops=dict(facecolor="white", edgecolor="#333333", linewidth=1.25),
    whiskerprops=dict(color="#333333", linewidth=1.15),
    capprops=dict(color="#333333", linewidth=1.15),
    zorder=2,
)
ax.axhline(0, color="#555555", linestyle=(0, (4, 3)), linewidth=1.3, zorder=1)

rng = np.random.default_rng(20260820)
for xpos, metric in zip(positions, metric_order_f):
    metric_rows = [r for r in rows_f if r["Pollutant"] == metric]
    jitter = rng.uniform(-0.15, 0.15, size=len(metric_rows))
    for j, row in enumerate(metric_rows):
        st = str(row["Station type"])
        style = TYPE_STYLE[st]
        ax.scatter(
            xpos + jitter[j],
            float(row["Offset Δt (months) [Y]"]),
            s=58,
            marker=style["marker"],
            color=style["color"],
            edgecolors="white",
            linewidths=0.7,
            zorder=4,
        )

    summary = summary_f[metric]
    n = int(summary["n"])
    med_abs = float(summary["Median |Δt| (months)"])
    ax.text(
        xpos,
        94,
        f"n = {n}",
        ha="center",
        va="top",
        fontsize=9.7,
        fontweight="bold",
        bbox=dict(facecolor="white", edgecolor="none", alpha=0.94, pad=0.5),
        zorder=7,
    )
    ax.text(
        xpos,
        85,
        f"med |Δt| = {med_abs:g}",
        ha="center",
        va="top",
        fontsize=8.8,
        fontweight="bold",
        bbox=dict(facecolor="white", edgecolor="none", alpha=0.94, pad=0.45),
        zorder=7,
    )

ax.set_xticks(positions)
ax.set_xticklabels([METRIC_DISPLAY[m] for m in metric_order_f], fontsize=11.5, fontweight="bold")
ax.set_xlim(-0.55, len(metric_order_f) - 0.45)
ax.set_ylim(-85, 100)
ax.set_yticks([-75, -50, -25, 0, 25, 50, 75])
ax.set_ylabel("Breakpoint offset relative to LDSA (months)", fontsize=13.5, fontweight="bold", labelpad=7)
ax.set_xlabel(
    "Negative = metric earlier; positive = metric later",
    fontsize=12.5,
    fontweight="bold",
    labelpad=10,
)
panel_heading(ax, "f", "Metric-specific transition timing")
style_axes(ax)
ax.grid(axis="y", color="#D9D9D9", linestyle=(0, (1.4, 2.2)), linewidth=0.8)
ax.grid(axis="x", visible=False)

handles_type = [
    Line2D(
        [0],
        [0],
        marker=TYPE_STYLE[st]["marker"],
        color="none",
        markerfacecolor=TYPE_STYLE[st]["color"],
        markeredgecolor="white",
        markeredgewidth=0.7,
        markersize=8,
        label=st,
    )
    for st in ["UB", "SUB", "TR", "RB"]
]
leg_f = ax.legend(
    handles=handles_type,
    loc="upper center",
    bbox_to_anchor=(0.5, -0.14),
    ncol=4,
    frameon=False,
    fontsize=10.5,
    columnspacing=2.2,
    handletextpad=0.5,
)
for text in leg_f.get_texts():
    text.set_fontweight("bold")


# Apply consistent typography to every visible text artist, including
# tick labels, annotations, statistics, legends, pollutant names and units.
for text_artist in fig.findobj(mpl.text.Text):
    text_artist.set_fontfamily("Times New Roman")
    text_artist.set_fontweight("bold")
    text_artist.set_fontsize(text_artist.get_fontsize() + 4)


base = OUTPUT_DIR / "Fig3_all_panels_python_redraw_labeled"
fig.savefig(base.with_suffix(".png"), dpi=500, bbox_inches="tight", facecolor="white")
fig.savefig(base.with_suffix(".pdf"), bbox_inches="tight", facecolor="white")
fig.savefig(base.with_suffix(".svg"), bbox_inches="tight", facecolor="white")
plt.close(fig)

print(base.with_suffix(".png"))
print(base.with_suffix(".pdf"))
print(base.with_suffix(".svg"))
