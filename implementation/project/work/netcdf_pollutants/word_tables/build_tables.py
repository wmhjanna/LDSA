from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
import sys, json, ast, hashlib, re

ROOT = Path('@PROJECT_ROOT@')
sys.path.insert(0, str(ROOT / 'work/ldsa_support/deps'))
import numpy as np
import pandas as pd
from scipy.stats import pearsonr
import ruptures as rpt
from lxml import etree as E
from docx import Document

W = Path(__file__).parent
A = ROOT / 'outputs/netcdf_pollutants/analysis'
A65 = ROOT / 'outputs/quality_control/analysis'
BASE = Path('@DELIVERY_ROOT@')
DEST = BASE / 'netcdf_pollutant_figures/tables'
SRC = BASE / 'provider_quality_control_analysis/tables/LDSA_supplementary_tables_S1-S6_quality_control.docx'
OUT = DEST / 'supplementary_pollutant_tables.docx'
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
qn = lambda s: '{' + ns['w'] + '}' + s
tree = ast.parse((ROOT / 'work/table_formatting/build_tables.py').read_text())
for f in tree.body:
    if isinstance(f, ast.FunctionDef) and f.name in ['rprop', 'set_p', 'set_cell', 'fmt', 'caption']:
        exec(compile(ast.Module(body=[f], type_ignores=[]), 'reference_cell_helpers', 'exec'))

def pval(value):
    return '—' if not np.isfinite(value) else '<0.001' if value < .001 else f'{value:.3f}'

def year_range(first, last):
    y0, y1 = pd.Timestamp(first).year, pd.Timestamp(last).year
    return str(y0) if y0 == y1 else f'{y0}–{y1}'

def iso_country(value):
    countries = {'Germany': 'DE', 'Finland': 'FI', 'France': 'FR', 'Sweden': 'SE',
                 'Norway': 'NO', 'Switzerland': 'CH', 'Italy': 'IT', 'Spain': 'ES',
                 'Czech Republic': 'CZ', 'Czechia': 'CZ', 'Greece': 'GR',
                 'Hungary': 'HU', 'United Kingdom': 'GB', 'UK': 'GB',
                 'Ireland': 'IE', 'Netherlands': 'NL', 'Belgium': 'BE',
                 'Bulgaria': 'BG', 'Malta': 'MT', 'Lithuania': 'LT',
                 'United States': 'US', 'USA': 'US', 'United States of America': 'US'}
    out = countries.get(str(value), str(value))
    assert len(out) == 2, ('Unmapped country', value)
    return out

inputs = [A / f for f in ['trends.csv', 'monthly_series.csv', 'full_monthly_series.csv',
                           'site_eligibility.csv', 'exact_overlap_pairs.csv',
                           'exact_overlap_breakpoints_min12.csv']]
inputs += [A65 / 'site_spatial_statistics.csv']
assert '--confirmed-ready' in sys.argv, 'Await explicit confirmation that all new analysis inputs are final'
for path in inputs:
    assert path.exists(), path
DEST.mkdir(parents=True, exist_ok=True)
with ZipFile(SRC) as z:
    parts = {i.filename: z.read(i.filename) for i in z.infolist()}
root = E.fromstring(parts['word/document.xml'])
body = root.find('w:body', ns)
tables = body.findall('w:tbl', ns)
assert len(tables) == 6
original = Document(SRC)
original_s2 = E.tostring(tables[1])
original_s6 = E.tostring(tables[5])
original_s4_internal = [E.tostring(row) for row in tables[3].findall('w:tr', ns)[1:4]]
original_section = E.tostring(body.find('w:sectPr', ns))
reference_rows = {f'S{i+1}': [[c.text for c in r.cells] for r in t.rows] for i, t in enumerate(original.tables)}
matrices = {}
sp = pd.read_csv(A65 / 'site_spatial_statistics.csv').set_index('Site')
assert len(sp) == 65 and sp.index.is_unique
mapped = sorted(sp.index, key=lambda s: ({'RB': 0, 'SUB': 1, 'TR': 2, 'UB': 3}[sp.loc[s, 'Type']], s))
tr_all = pd.read_csv(A / 'trends.csv')
tr = tr_all[~tr_all.Site.isin(['PRA_SUB', 'PAR_SUB'])].copy()
ms = pd.read_csv(A / 'monthly_series.csv', parse_dates=['date'])
ms = ms[~ms.Site.isin(['PRA_SUB', 'PAR_SUB'])].copy()
fm = pd.read_csv(A / 'full_monthly_series.csv', parse_dates=['date'])
internal_bp_path = ROOT / 'outputs/quality_control/workbook/analysis/breakpoints.csv'
bp = pd.read_csv(internal_bp_path)
inputs.append(internal_bp_path)
pairs = pd.read_csv(A / 'exact_overlap_pairs.csv')
pairs = pairs[~pairs.Site.isin(['PRA_SUB', 'PAR_SUB'])].copy()
eligibility = pd.read_csv(A / 'site_eligibility.csv')
metrics = ['PM2.5', 'PM10', 'BC', 'NO2', 'SO2', 'NO', 'CO', 'O3']
polorder = ['PM1', 'PM2.5', 'PM10', 'BC', 'SO2', 'NO', 'NO2', 'O3', 'CO']
assert not ((tr.Site == 'HEL_TR') & (tr.Metric == 'SO2')).any()
assert not ((fm.Site == 'HEL_TR') & (fm.Metric == 'SO2')).any()

# Only the ancillary column changes; the current 65-site metadata stays intact.
rows = []
for row, old in zip(tables[0].findall('w:tr', ns)[1:], reference_rows['S1'][1:]):
    values = list(old)
    available = set(fm.loc[fm.Site.eq(values[0]) & fm.value.notna(), 'Metric'])
    values[7] = ', '.join(metric for metric in polorder if metric in available) or 'Not available'
    set_cell(row.findall('w:tc', ns)[7], values[7], chemical=True)
    rows.append(values)
matrices['S1'] = rows
pd.DataFrame(rows, columns=['Site', 'Country', 'Class', 'Measurement period', 'Spatial period',
                           'Period classification', 'Native size ranges (nm)', 'Available co-pollutants']).to_csv(DEST / 'Table_S1_network.csv', index=False)
matrices['S2'] = reference_rows['S2'][1:]
# S3: correlations and trend classifications.
rows, raw = [], []
for metric in metrics:
    for strategy, inc in [('Full available record', None), ('Exact overlap ≥60 months (primary)', 'primary_included'), ('Exact overlap ≥36 months (sensitivity)', 'sensitivity_included')]:
        if inc is None:
            g = tr[tr.Phase.eq('Full') & tr.Metric.isin(['LDSA', metric])].pivot(index='Site', columns='Metric', values='Slope_Percent').replace([np.inf, -np.inf], np.nan).dropna()
            x, y = g[metric], g.LDSA
            discord, dn, dd = '—', None, None
        else:
            g = pairs[pairs.Pollutant.eq(metric) & pairs[inc]].copy()
            g = g[np.isfinite(g['Pollutant slope (% yr^-1) [X]']) & np.isfinite(g['LDSA slope (% yr^-1) [Y]'])]
            x, y = g['Pollutant slope (% yr^-1) [X]'], g['LDSA slope (% yr^-1) [Y]']
            decl = g.Pollutant_class.eq('Declining')
            dd, dn = int(decl.sum()), int((decl & g.LDSA_class.ne('Declining')).sum())
            discord = f'{dn}/{dd} ({100*dn/dd:.1f}%)' if dd else '0/0 (—)'
        cor = pearsonr(x, y) if len(g) > 2 and np.std(x) > 0 and np.std(y) > 0 else (np.nan, np.nan)
        rows.append([metric, strategy, str(len(g)), fmt(cor[0], 3), pval(cor[1]), discord])
        raw.append(dict(Metric=metric, Strategy=strategy, n=len(g), r=cor[0], p=cor[1], discordant_n=dn, pollutant_declining_n=dd))
matrices['S3'] = rows
pd.DataFrame(raw).to_csv(DEST / 'Table_S3_values.csv', index=False)
for row, values in zip(tables[2].findall('w:tr', ns)[1:], rows):
    for j in range(2, 6):
        set_cell(row.findall('w:tc', ns)[j], values[j])

# S4 conventional monthly segmentation reproduces Figure 3 and includes PM1.
timing = []
for metric in ['PM1'] + metrics:
    for site, g in ms[ms.Metric.isin(['LDSA', metric])].groupby('Site'):
        q = g.pivot(index='date', columns='Metric', values='value')
        if not {'LDSA', metric}.issubset(q):
            continue
        q = q.dropna().sort_index()
        n = len(q)
        if n < 60:
            continue
        span = (q.index[-1].year - q.index[0].year) * 12 + q.index[-1].month - q.index[0].month + 1
        if n / span < .5:
            continue
        dates = {}
        for col in ['LDSA', metric]:
            y = q[col].to_numpy()
            z = (y - y.mean()) / y.std()
            k = rpt.KernelCPD(kernel='linear', min_size=12, jump=1).fit(z.reshape(-1, 1)).predict(n_bkps=1)[0]
            dates[col] = q.index[k]
        offset = (dates[metric].year - dates['LDSA'].year) * 12 + dates[metric].month - dates['LDSA'].month
        timing.append(dict(Site=site, Metric=metric, offset=offset, n_months=n, coverage=n/span))
timing = pd.DataFrame(timing)
timing.to_csv(DEST / 'Table_S4_conventional_site_pairs.csv', index=False)
figtim = pd.read_csv(A / 'exact_overlap_breakpoints_min12.csv')
figtim = figtim[~figtim.Site.isin(['PAR_SUB', 'PRA_SUB'])]
check = timing[timing.Metric.ne('PM1')].merge(figtim, left_on=['Site', 'Metric'], right_on=['Site', 'Pollutant'])
assert len(check) == len(figtim) == len(timing[timing.Metric.ne('PM1')])
assert np.array_equal(check.offset, check['Offset Δt (months) [Y]'])
rows, raw = [], []
for old in reference_rows['S4'][1:]:
    group, metric = old[:2]
    v = bp.loc[bp.Metric.eq(metric), 'offset_months'].dropna().to_numpy() if group == 'PNSD-derived' else timing.loc[timing.Metric.eq(metric), 'offset'].to_numpy()
    n = len(v)
    assert n > 0, (group, metric)
    q1, q3 = np.quantile(v, [.25, .75])
    counts = [int((np.abs(v) <= 12).sum()), int((np.abs(v) > 24).sum()), int((v < 0).sum()), int((v > 0).sum())]
    values = [group, metric, str(n), fmt(np.median(v), 1), f'{fmt(q1,2).rstrip("0").rstrip(".")} to {fmt(q3,2).rstrip("0").rstrip(".")}', fmt(np.median(np.abs(v)), 1)] + [f'{x} ({100*x/n:.1f}%)' for x in counts]
    rows.append(values)
    raw.append(dict(group=group, Metric=metric, n=n, median=np.median(v), q1=q1, q3=q3,
                    median_absolute=np.median(np.abs(v)), within12=counts[0], over24=counts[1], earlier=counts[2], later=counts[3]))
matrices['S4'] = rows
pd.DataFrame(raw).to_csv(DEST / 'Table_S4_values.csv', index=False)
for row, values in zip(tables[3].findall('w:tr', ns)[1:], rows):
    for j in range(2, 10):
        set_cell(row.findall('w:tc', ns)[j], values[j])

# S5 phase statistics use exactly the same excluded-site population as Figure S8.
rows, raw = [], []
for old in reference_rows['S5'][1:]:
    metric = old[0]
    values = [metric]
    for phase in ['Before', 'After']:
        g = tr[tr.Metric.eq(metric) & tr.Phase.eq(phase)]
        g = g[np.isfinite(g.Slope_Percent)]
        median = g.Slope_Percent.median()
        values.extend([str(len(g)), fmt(median, 2, True)])
        raw.append(dict(Metric=metric, Phase=phase, n=len(g), median=median))
    rows.append(values)
matrices['S5'] = rows
pd.DataFrame(raw).to_csv(DEST / 'Table_S5_values.csv', index=False)
for row, values in zip(tables[4].findall('w:tr', ns)[1:], rows):
    for j in range(1, 5):
        set_cell(row.findall('w:tc', ns)[j], values[j])
matrices['S6'] = reference_rows['S6'][1:]

caption(tables[0].getprevious(), 'Table S1. Observational network, particle-size coverage, spatial-analysis periods and available ancillary measurements.',
        'The network includes 65 physical sites. Spatial statistics use available observations in 2015–2019 or, when that window has no valid observations, the nearest available year. Measurement periods span valid primary LDSA observations. Size ranges show native diameter centers across instrument configurations. RB, SUB, TR and UB denote regional background, suburban background, traffic and urban background. Co-pollutant entries describe available monthly analysis inputs across the network; Not available means no co-pollutant monthly series was used. Paris and Prague each contribute one combined LDSA record.')
caption(tables[2].getprevious(), 'Table S3. Sensitivity of LDSA–conventional-pollutant trend correspondence to temporal alignment.',
        'Pearson correlations compare site-level Theil–Sen slopes. Exact-overlap analyses use identical valid months, with at least 14 valid days per monthly mean and at least 50% calendar coverage. Discordance denotes no statistically supported LDSA decline among records with a statistically supported pollutant decline. Full-record results use each metric’s available months. Comparisons include eligible co-pollutant records from the study network, excluding PRA_SUB and PAR_SUB; HEL_TR SO₂ is excluded.')
caption(tables[3].getprevious(), 'Table S4. Timing agreement between LDSA and PNSD-derived or conventional aerosol-metric breakpoints.',
        'Offsets are Δt = tmetric − tLDSA, in calendar months; negative values indicate earlier metric transitions. PNSD-derived modes are particle-size-structure diagnostics from the observation sequences from workbook-based PNSD inputs. Conventional pollutants use primary exact-overlap monthly analyses across the study network, excluding PRA_SUB and PAR_SUB. Counts refer to physical sites.')
bp_count = eligibility[~eligibility.Site.isin(['PAR_SUB', 'PRA_SUB']) & eligibility.breakpoint.notna()].Site.nunique()
caption(tables[4].getprevious(), 'Table S5. Phase-specific LDSA-anchored slopes of co-pollutants.',
        f'Median percentage slopes use phases defined by the primary LDSA breakpoint. Across the study network and after excluding PRA_SUB and PAR_SUB, breakpoints are available for {bp_count} physical sites; each phase requires at least three monthly observations for fitting. Co-pollutants are split at the start of the LDSA breakpoint month before direct monthly aggregation; the original LAN-specific script uses the exact breakpoint time. Sample sizes count independent records and vary by pollutant and phase. HEL_TR SO₂ is excluded.')

assert E.tostring(tables[1]) == original_s2
assert E.tostring(tables[5]) == original_s6
assert [E.tostring(row) for row in tables[3].findall('w:tr', ns)[1:4]] == original_s4_internal
assert E.tostring(body.find('w:sectPr', ns)) == original_section
parts['word/document.xml'] = E.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
with ZipFile(OUT, 'w', ZIP_DEFLATED) as z:
    for name, data in parts.items():
        z.writestr(name, data)

final = Document(OUT)
assert len(final.tables) == 6
for i, table in enumerate(final.tables):
    values = matrices[f'S{i+1}']
    assert len(table.rows) - 1 == len(values)
    for row, vals in zip(table.rows[1:], values):
        for j, val in enumerate(vals):
            if i == 2 and j in [0, 1]:
                continue
            assert row.cells[j].text == val, (i, j, row.cells[j].text, val)
with ZipFile(OUT) as z, ZipFile(SRC) as orig:
    changed = [name for name in parts if z.read(name) != orig.read(name)]
    assert changed == ['word/document.xml']
    assert z.testzip() is None
(W / 'table_values.json').write_text(json.dumps(matrices, ensure_ascii=False, indent=2))
changes = []
for key, current in matrices.items():
    oldrows = reference_rows[key][1:]
    if key in ['S1', 'S6']:
        lookup = {row[0]: row for row in oldrows}
        matched = [(row[0], lookup.get(row[0]), row) for row in current]
    else:
        matched = [(str(i + 1), old, new) for i, (old, new) in enumerate(zip(oldrows, current))]
    for row_id, old, new in matched:
        for j, val in enumerate(new):
            before = None if old is None else old[j]
            if before != val:
                changes.append(dict(Table=key, Row=row_id, Column=j + 1, Before=before, After=val,
                                    Change='Added site' if old is None else 'Updated'))
pd.DataFrame(changes).to_csv(DEST / 'Table_S1-S6_changes.csv', index=False)
verification = dict(output=str(OUT), source=str(SRC), source_sha256=hashlib.sha256(SRC.read_bytes()).hexdigest(),
                    output_sha256=hashlib.sha256(OUT.read_bytes()).hexdigest(),
                    changed_package_parts=changed, unchanged_S2=True, unchanged_S6=True, unchanged_S4_internal=True, section_preserved=True,
                    all_table_cells_verified=True, table_data_rows=[len(matrices[f'S{i}']) for i in range(1, 7)],
                    S4_matches_figure3=len(check), co_pollutant_excluded=['PAR_SUB', 'PRA_SUB'],
                    HEL_TR_SO2_excluded=True,
                    inputs=[dict(path=str(p), sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in inputs])
(W / 'structural_verification.json').write_text(json.dumps(verification, indent=2))
print(json.dumps(verification, ensure_ascii=False, indent=2))
